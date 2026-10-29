"""pcf_format.py — Module canonique de mise en forme des factures.

Source unique du format : ce module + le gabarit `facturation/assets/` > MODELE.
Toute spec visuelle ailleurs est obsolète (cf. references/FORMAT.md).

Usage : voir references/PROC00_pcf.md étape 5.2.

Gabarit (ADR015) : une mise en page unique, `PCF_facture_standard.xlsx`, vide de
toute coordonnée ; seules les informations de l'émetteur changent d'une entité à
l'autre. Une entité peut avoir son propre gabarit `PCF_facture_<ENTITE>.xlsx`,
prioritaire s'il existe.

Source des données :
- Coordonnées émetteur, destinataires : `MASTER.Coordonnées` (Type=Émetteur/Destinataire)
- SIRET, IBAN, BIC, TVA : `MASTER.Coordonnées` (Type=Info, Code=<ENTITE>)
- Aucune duplication de données dans ce module — tout est lu depuis MASTER au runtime.
  Le gabarit versionné est volontairement vide de toute coordonnée : elles sont
  injectées par apply_emetteur / apply_metadata / apply_rib.

Où le module cherche les données (racine de données) :
1. `set_data_root(chemin)` si appelé explicitement ;
2. sinon la variable d'environnement `PCF_DATA_ROOT` ;
3. sinon le répertoire de travail courant.
La racine de données doit contenir `<ENTITE>/<ENTITE>_MASTER.xlsx`.
"""

import os
import re
from copy import copy as _copy
from datetime import date
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

# ──────────────────────────────────────────────────────────────────────
# Chemins
# ──────────────────────────────────────────────────────────────────────

ASSETS = Path(__file__).resolve().parent / "assets"
GABARIT_STANDARD = ASSETS / "PCF_facture_standard.xlsx"

# Code d'entité : majuscules et chiffres, 2 à 16 caractères (ex. ROSARI, KAUR).
_ENTITE_RE = re.compile(r"^[A-Z][A-Z0-9]{1,15}$")

_DATA_ROOT: Optional[Path] = None


def valider_entite(entite: str) -> str:
    """Vérifie le format d'un code d'entité et le renvoie.

    Raises:
        ValueError: si le code n'est pas en majuscules (2 à 16 caractères) ou
            s'il vaut STANDARD, réservé au gabarit commun.
    """
    if not isinstance(entite, str) or not _ENTITE_RE.match(entite):
        raise ValueError(
            f"Code d'entité invalide : {entite!r}. Attendu : majuscules et "
            "chiffres, 2 à 16 caractères (ex. ROSARI, KAUR)."
        )
    if entite == "STANDARD":
        raise ValueError("Le code d'entité STANDARD est réservé au gabarit commun.")
    return entite


def gabarit_pour(entite: str) -> Path:
    """Gabarit à utiliser pour une entité : le sien s'il existe, sinon le standard."""
    propre = ASSETS / f"PCF_facture_{valider_entite(entite)}.xlsx"
    return propre if propre.exists() else GABARIT_STANDARD


def set_data_root(chemin) -> Path:
    """Fixe la racine des données métier pour la session Python en cours.

    Args:
        chemin: dossier contenant `<ENTITE>/<ENTITE>_MASTER.xlsx`.

    Returns:
        Le chemin résolu.
    """
    global _DATA_ROOT
    _DATA_ROOT = Path(chemin).expanduser().resolve()
    _MASTER_DATA_CACHE.clear()
    return _DATA_ROOT


def get_data_root() -> Path:
    """Racine des données : set_data_root() > $PCF_DATA_ROOT > dossier courant."""
    if _DATA_ROOT is not None:
        return _DATA_ROOT
    env = os.environ.get("PCF_DATA_ROOT")
    if env:
        return Path(env).expanduser().resolve()
    return Path.cwd()

# ──────────────────────────────────────────────────────────────────────
# Cache lecture MASTER.Coordonnées (lazy loading)
# ──────────────────────────────────────────────────────────────────────

_MASTER_DATA_CACHE: dict = {}


def _master_path_for(entite: str) -> Path:
    """Retourne le chemin du MASTER de l'entité, sous la racine de données."""
    valider_entite(entite)
    return get_data_root() / entite / f"{entite}_MASTER.xlsx"


def _load_master_data(entite: str, refresh: bool = False) -> dict:
    """Charge MASTER.Coordonnées et retourne un dict structuré.

    Format :
    {
        "emetteur": {"nom": ..., "adresse_l1": ..., "adresse_l2": ..., "email": ..., "tel": ...},
        "destinataires": {
            "DOR": {...},
            "EPITA": {...},
            ...
        },
        "info": {"SIRET": ..., "IBAN": ..., "BIC": ..., "TVA": ...},
    }

    Args:
        entite: "ROSARI" ou "KAUR".
        refresh: si True, force la relecture du MASTER (ignore le cache).
    """
    if not refresh and entite in _MASTER_DATA_CACHE:
        return _MASTER_DATA_CACHE[entite]

    master_path = _master_path_for(entite)
    if not master_path.exists():
        raise FileNotFoundError(
            f"MASTER introuvable : {master_path}\n"
            f"Racine de données utilisée : {get_data_root()}\n"
            "Corrige-la via set_data_root(...) ou la variable d'env PCF_DATA_ROOT."
        )

    wb = load_workbook(master_path, data_only=True)
    if "Coordonnées" not in wb.sheetnames:
        raise RuntimeError(
            f"Onglet 'Coordonnées' manquant dans {master_path.name}"
        )
    ws = wb["Coordonnées"]

    # Header attendu : Type | Code | Nom | Adresse L1 | Adresse L2 | Adresse L3 | Email | Téléphone
    data = {"emetteur": None, "destinataires": {}, "info": {}}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        type_ = row[0]
        code = row[1]
        nom = row[2]
        adr_l1 = row[3]
        adr_l2 = row[4]
        adr_l3 = row[5]
        email = row[6]
        tel = row[7]

        if type_ == "Émetteur" and code == entite:
            data["emetteur"] = {
                "nom": nom or "",
                "adresse_l1": adr_l1 or "",
                "adresse_l2": adr_l2 or "",
                "email": email or "",
                "tel": tel or "",
            }
        elif type_ == "Destinataire" and code:
            # Stocker None si toutes les valeurs sont vides
            if not nom and not adr_l1 and not adr_l2 and not email:
                data["destinataires"][code] = None
            else:
                data["destinataires"][code] = {
                    "nom": nom or "",
                    "adresse_l1": adr_l1 or "",
                    "adresse_l2": adr_l2 or "",
                    "adresse_l3": adr_l3 or "",
                    "email": email or "",
                }
        elif type_ == "Info" and code == entite:
            # Info row : champ Nom contient le label (SIRET/IBAN/BIC/TVA)
            # et Adresse L1 contient la valeur
            if nom:
                data["info"][nom] = adr_l1 or ""

    if data["emetteur"] is None:
        raise RuntimeError(
            f"Pas d'émetteur '{entite}' trouvé dans {master_path.name}.Coordonnées"
        )
    required_info = {"SIRET", "IBAN", "BIC"}
    missing = required_info - set(data["info"].keys())
    if missing:
        raise RuntimeError(
            f"Info manquantes pour {entite} dans MASTER.Coordonnées : {missing}. "
            f"Attendu lignes : Type=Info, Code={entite}, Nom=SIRET/IBAN/BIC/TVA."
        )

    _MASTER_DATA_CACHE[entite] = data
    return data


def get_destinataire(code_ecole: str, entite: str = "ROSARI") -> Optional[dict]:
    """Helper public : retourne le destinataire pour une école donnée, ou None.

    Args:
        code_ecole: code école (ex. "DOR", "EPITA").
        entite: "ROSARI" ou "KAUR".
    """
    data = _load_master_data(entite)
    return data["destinataires"].get(code_ecole)


# Plan des cellules du gabarit standard, commun à toutes les entités (ADR015).
LAYOUT_STANDARD = {
    "bandeau_row": 2,
    "emetteur_rows": (7, 11),
    "emetteur_col": 3,  # C
    "destinataire_rows": (7, 10),
    "destinataire_col": 13,  # M
    "bon_commande_row": 14,
    "bon_commande_label_col": 3,  # C
    "bon_commande_value_cols": (4, 8),  # D:H (merge)
    "siret_label_cell": "C16",
    "siret_cell": "D16",
    "n_facture_label_cell": "C17",
    "n_facture_cell": "D17",
    "table_header_row": 19,
    "cours_first_row": 20,
    "cours_last_row": 79,  # 60 rows
    "cols_table": {
        "ecole": 3,    # C (merge C:D)
        "promo": 5,    # E (merge E:F)
        "libelle": 7,  # G (merge G:J)
        "date": 11,    # K (merge K:L)
        "heures": 13,  # M (merge M:N)
        "taux": 15,    # O (merge O:P)
        "total": 17,   # Q (merge Q:R)
    },
    "total_row": 80,
    "rib_iban_cell": "D88",
    "rib_bic_cell": "D89",
    "derniere_ligne_utile": 93,  # mentions légales → borne du print_area
}

LAYOUT_ROSARI = LAYOUT_STANDARD  # nom historique, conservé pour les scripts existants

# Plans propres à une entité dont le gabarit diffère du standard (aucun pour l'instant).
LAYOUTS: dict = {}

THIN = Side(style="thin", color="000000")

# ──────────────────────────────────────────────────────────────────────
# Helpers internes
# ──────────────────────────────────────────────────────────────────────


def _set_if_not_merged(ws: Worksheet, row: int, col: int, value):
    """Écrit value si la cell n'est pas une MergedCell continuation."""
    cell = ws.cell(row=row, column=col)
    if not isinstance(cell, MergedCell):
        cell.value = value


def _unmerge_in_range(
    ws: Worksheet, rmin: int, rmax: int, cmin: int = 1, cmax: int = 30
):
    """Démerge toutes les plages incluses dans la zone."""
    to_remove = []
    for mr in list(ws.merged_cells.ranges):
        if (
            mr.min_row >= rmin
            and mr.max_row <= rmax
            and mr.min_col >= cmin
            and mr.max_col <= cmax
        ):
            to_remove.append(str(mr))
    for mr_str in to_remove:
        try:
            ws.unmerge_cells(mr_str)
        except Exception:
            pass


def _frame_box(ws: Worksheet, rmin: int, rmax: int, cmin: int, cmax: int):
    """Trace un encadré propre : cadre extérieur fin, aucun trait horizontal interne.

    Le template DOR001 d'origine porte des bordures résiduelles au milieu des
    encadrés émetteur/destinataire (la zone faisait 3 lignes avant d'être
    étendue à 5). Sans ce nettoyage, un trait traverse l'encadré.
    """
    # Note : les bordures sont assignables sur les MergedCell (contrairement
    # aux valeurs), il faut donc parcourir toute la zone et pas seulement les
    # cellules d'ancrage, sinon le cadre droit manque sur les zones fusionnées.
    for r in range(rmin, rmax + 1):
        for c in range(cmin, cmax + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = Border(
                top=THIN if r == rmin else None,
                bottom=THIN if r == rmax else None,
                left=THIN if c == cmin else None,
                right=THIN if c == cmax else None,
            )


def _xworkbook_clone_modele(
    template_path: Path, dst_wb: Workbook, new_name: str
) -> Worksheet:
    """Charge MODELE depuis template_path et le recopie cellule par cellule dans dst_wb.

    Le gabarit compte environ 26 800 cellules stylées mais une soixantaine de styles
    distincts. Chaque style source n'est donc recopié qu'une fois (police, remplissage,
    bordure, alignement, format de nombre) ; les cellules suivantes qui portent le même
    style reçoivent directement les index du style déjà créé dans dst_wb. Le résultat est
    identique à une copie attribut par attribut, en 20 à 50 fois moins de temps.
    """
    src_wb = load_workbook(template_path)
    src_ws = src_wb["MODELE"]

    new_ws = dst_wb.create_sheet(new_name)

    # Copie des cellules : un style source -> un style destination, recopié une seule fois.
    styles_recopies: dict[tuple, object] = {}
    for row in src_ws.iter_rows():
        for cell in row:
            if cell.value is None and not cell.has_style:
                continue
            new_cell = new_ws.cell(row=cell.row, column=cell.column)
            new_cell.value = cell.value
            if not cell.has_style:
                continue
            cle = tuple(cell._style)
            deja = styles_recopies.get(cle)
            if deja is not None:
                new_cell._style = _copy(deja)
                continue
            new_cell.font = _copy(cell.font)
            new_cell.fill = _copy(cell.fill)
            new_cell.border = _copy(cell.border)
            new_cell.alignment = _copy(cell.alignment)
            new_cell.number_format = cell.number_format
            styles_recopies[cle] = _copy(new_cell._style)

    # Column widths
    for col_letter, dim in src_ws.column_dimensions.items():
        if dim.width:
            new_ws.column_dimensions[col_letter].width = dim.width

    # Row heights
    for r, dim in src_ws.row_dimensions.items():
        if dim.height:
            new_ws.row_dimensions[r].height = dim.height

    # Merges
    for mr in src_ws.merged_cells.ranges:
        try:
            new_ws.merge_cells(str(mr))
        except Exception:
            pass

    return new_ws


def _layout_for(entite: str) -> dict:
    """Plan des cellules de l'entité : le sien s'il est déclaré, sinon le standard."""
    return LAYOUTS.get(valider_entite(entite), LAYOUT_STANDARD)


def _emetteur_for(entite: str) -> dict:
    """Retourne l'émetteur lu depuis MASTER.Coordonnées."""
    return _load_master_data(entite)["emetteur"]


def _rib_for(entite: str) -> dict:
    """Retourne {IBAN, BIC} lu depuis MASTER.Coordonnées (Type=Info)."""
    info = _load_master_data(entite)["info"]
    return {"IBAN": info["IBAN"], "BIC": info["BIC"]}


def _siret_for(entite: str) -> str:
    """Retourne le SIRET lu depuis MASTER.Coordonnées (Type=Info)."""
    return _load_master_data(entite)["info"]["SIRET"]


# ──────────────────────────────────────────────────────────────────────
# API publique
# ──────────────────────────────────────────────────────────────────────


def clone_template(
    dst_wb: Workbook, sheet_name: str, entite: str = "ROSARI"
) -> Worksheet:
    """Charge MODELE depuis le template de l'entité et le copie dans dst_wb.

    Args:
        dst_wb: workbook de destination (doit déjà exister, peut être vide ou avoir d'autres sheets).
        sheet_name: nom de la nouvelle feuille (ex. "R_DOR001 Mars 2026").
        entite: code de l'entité (ex. "ROSARI").

    Returns:
        Le nouveau worksheet, prêt à recevoir les données via les autres fonctions.
    """
    gabarit = gabarit_pour(entite)
    if not gabarit.exists():
        raise FileNotFoundError(
            f"Gabarit introuvable : {gabarit}. "
            "Le régénérer avec scripts/build_template.py."
        )
    return _xworkbook_clone_modele(gabarit, dst_wb, sheet_name)


def apply_emetteur(ws: Worksheet, entite: str = "ROSARI"):
    """Re-écrit l'émetteur dans le bloc col C (ROSARI) ou équivalent.

    Idempotent : peut être appelé sur un template déjà rempli.

    Args:
        ws: worksheet cible (créé via clone_template).
        entite: "ROSARI" ou "KAUR".
    """
    layout = _layout_for(entite)
    emetteur = _emetteur_for(entite)
    rmin, rmax = layout["emetteur_rows"]
    col = layout["emetteur_col"]

    lignes = [
        emetteur["nom"],
        emetteur["adresse_l1"],
        emetteur["adresse_l2"],
        emetteur["email"],
        emetteur["tel"],
    ]
    for i, valeur in enumerate(lignes):
        r = rmin + i
        _set_if_not_merged(ws, r, col, valeur)
        cell = ws.cell(row=r, column=col)
        if not isinstance(cell, MergedCell):
            if i == 0:
                cell.font = Font(bold=True, size=12)
            cell.alignment = Alignment(horizontal="center", vertical="center")

    _frame_box(ws, rmin, rmax, col, col + 2)


def apply_destinataire(
    ws: Worksheet, dest: Optional[dict], entite: str = "ROSARI"
):
    """Écrit le destinataire ou vide la zone si dest=None.

    Si dest=None : vide les valeurs (l'encadré reste, c'est un placeholder).
    Si dest=dict : écrit nom, adresse_l1, adresse_l2 (+ l3), email.

    Le template fusionne M7:P8 (le nom sur 2 lignes), ce qui faisait tomber la
    1re ligne d'adresse dans une MergedCell et la supprimait silencieusement.
    On démerge donc la zone et on refusionne une plage par ligne, pour disposer
    d'autant de slots que de lignes d'adresse.

    Args:
        ws: worksheet cible.
        dest: dict avec keys nom/adresse_l1/adresse_l2/adresse_l3/email, ou None.
        entite: "ROSARI" ou "KAUR".

    dest dict structure :
        {"nom": str, "adresse_l1": str, "adresse_l2": str,
         "adresse_l3": str (optional), "email": str (optional)}
    """
    layout = _layout_for(entite)
    rmin, rmax = layout["destinataire_rows"]
    col = layout["destinataire_col"]
    col_end = col + 3  # M:P

    # Normaliser la zone : une plage fusionnée par ligne
    _unmerge_in_range(ws, rmin, rmax, col, col_end)
    for r in range(rmin, rmax + 1):
        ws.merge_cells(
            start_row=r, start_column=col, end_row=r, end_column=col_end
        )
        ws.cell(row=r, column=col).value = None

    _frame_box(ws, rmin, rmax, col, col_end)

    if dest is None:
        # Encadré vide conservé comme placeholder
        return

    # Ville et pays sur la même ligne : la zone n'offre que rmax-rmin+1 slots
    ville = dest.get("adresse_l2", "")
    pays = dest.get("adresse_l3", "")
    if pays:
        ville = f"{ville}, {pays}" if ville else pays

    lignes = [
        v
        for v in (
            dest.get("nom", ""),
            dest.get("adresse_l1", ""),
            ville,
            dest.get("email", ""),
        )
        if v
    ]
    if len(lignes) > rmax - rmin + 1:
        raise RuntimeError(
            f"Destinataire à {len(lignes)} lignes pour une zone de "
            f"{rmax - rmin + 1} lignes (rows {rmin}-{rmax})."
        )

    for i, valeur in enumerate(lignes):
        r = rmin + i
        cell = ws.cell(row=r, column=col)
        cell.value = valeur
        cell.font = Font(bold=True, size=12) if i == 0 else Font(size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")


def apply_metadata(
    ws: Worksheet,
    n_facture: str,
    date_emission: date,
    entite: str = "ROSARI",
    ref_bon_commande: Optional[str] = None,
):
    """Écrit titre row 2 + Bon de commande row 14 + SIRET D16 + N° Facture D17.

    Args:
        ws: worksheet cible.
        n_facture: ex. "R_DOR001".
        date_emission: date d'émission de la facture.
        entite: "ROSARI" ou "KAUR".
        ref_bon_commande: réf. du bon de commande client (ex. "PO-26-ASSOTA-01791").
            Si None, la ligne n'est pas écrite. Obligatoire pour EPITA, dont la
            compta (Basware) rapproche facture et PO.
    """
    layout = _layout_for(entite)
    siret = _siret_for(entite)

    titre = f"FACTURE\nDate d'émission: {date_emission.strftime('%d/%m/%Y')}"
    _set_if_not_merged(ws, layout["bandeau_row"], 2, titre)

    ws[layout["siret_label_cell"]] = "N° SIRET :"
    ws[layout["siret_cell"]] = siret

    ws[layout["n_facture_label_cell"]] = "N° de Facture :"
    ws[layout["n_facture_cell"]] = n_facture

    if ref_bon_commande:
        row = layout["bon_commande_row"]
        label_col = layout["bon_commande_label_col"]
        vmin, vmax = layout["bon_commande_value_cols"]

        modele = ws[layout["n_facture_label_cell"]]
        for col, valeur in ((label_col, "Bon de commande :"), (vmin, ref_bon_commande)):
            cell = ws.cell(row=row, column=col)
            cell.value = valeur
            cell.font = _copy(modele.font)
            cell.alignment = _copy(modele.alignment)

        _unmerge_in_range(ws, row, row, vmin, vmax)
        ws.merge_cells(
            start_row=row, start_column=vmin, end_row=row, end_column=vmax
        )
        # Hauteur auto : la ligne fait 14.25pt dans le template, illisible en 18pt
        ws.row_dimensions[row].height = None


def place_sessions(
    ws: Worksheet,
    ecole_name: str,
    groups: list,
    entite: str = "ROSARI",
    masquer_jour: bool = False,
):
    """Place sessions groupées avec merges Promotion/Libellé/Jour/Heures/Taux/Total.

    Args:
        ws: worksheet cible (template cloné).
        ecole_name: nom de l'école pour la colonne ECOLE (ex. "DOR (NEXA Digital School)").
        groups: list of (promotion: str, libelle: str, sessions: list).
            Chaque session = (date: date | datetime, heures: float, taux: float).
        entite: "ROSARI" ou "KAUR".
        masquer_jour: masque les colonnes Jour (K:L). À activer quand la facture
            agrège un volume horaire sans séance datée (ex. EPITA : 1 ligne par
            UE de bon de commande). La date reste saisie, simplement non imprimée.

    Comportement :
        - Démerge la zone cours du template
        - Vide les valeurs cours
        - Place chaque session : date, heures, taux, formule total
        - Applique merges : Promotion E:F par groupe, Libellé G:J par sous-groupe,
          ECOLE C:D global sur les rows utilisées
        - Si nb_sessions < 60 : masque les rows excédentaires (jamais delete)
        - Si nb_sessions > 60 : insert_rows pour étendre, avec démerge préalable
    """
    layout = _layout_for(entite)
    cols = layout["cols_table"]
    first_row = layout["cours_first_row"]
    last_row = layout["cours_last_row"]
    capacity = last_row - first_row + 1  # 60 rows

    # Aplatir les groupes
    flat = []
    libelle_subranges = []
    promo_ranges = []
    cur_row = first_row
    last_promo = None
    promo_start = None
    for promo, libelle, sessions in groups:
        libelle_start = cur_row
        for sess in sessions:
            d, heures, taux = sess
            flat.append((promo, libelle, d, heures, taux))
            cur_row += 1
        libelle_end = cur_row - 1
        libelle_subranges.append((libelle, libelle_start, libelle_end))
        if promo != last_promo:
            if last_promo is not None:
                promo_ranges.append(
                    (last_promo, promo_start, libelle_start - 1)
                )
            last_promo = promo
            promo_start = libelle_start
    if last_promo is not None:
        promo_ranges.append((last_promo, promo_start, cur_row - 1))

    n_sessions = len(flat)
    last_session_row = first_row + n_sessions - 1

    # Extension dynamique si > capacity
    if n_sessions > capacity:
        n_extra = n_sessions - capacity
        # Démerger les plages sous la zone cours (TOTAL etc.)
        _unmerge_in_range(ws, last_row + 1, ws.max_row, 1, 30)
        # Insérer extra rows juste avant TOTAL
        ws.insert_rows(last_row + 1, n_extra)
        # Mettre à jour last_row
        new_last_row = last_row + n_extra
        # Appliquer pattern cours sur les nouvelles rows
        for r in range(last_row + 1, new_last_row + 1):
            ws.merge_cells(start_row=r, start_column=cols["promo"], end_row=r, end_column=cols["promo"] + 1)
            ws.merge_cells(start_row=r, start_column=cols["libelle"], end_row=r, end_column=cols["libelle"] + 3)
            ws.merge_cells(start_row=r, start_column=cols["date"], end_row=r, end_column=cols["date"] + 1)
            ws.merge_cells(start_row=r, start_column=cols["heures"], end_row=r, end_column=cols["heures"] + 1)
            ws.merge_cells(start_row=r, start_column=cols["taux"], end_row=r, end_column=cols["taux"] + 1)
            ws.merge_cells(start_row=r, start_column=cols["total"], end_row=r, end_column=cols["total"] + 1)
            ws.row_dimensions[r].height = 46.5
        last_row = new_last_row
        # Mettre à jour les formules TOTAL et À payer
        total_row = layout["total_row"] + n_extra
        ws.cell(row=total_row, column=cols["total"]).value = (
            f"=SUM({chr(64 + cols['total'])}{first_row}:{chr(64 + cols['total'] + 1)}{last_row})"
        )

    # Démerger zone cours actuelle (rows first_row à last_row, cols 3-18)
    # Important : démerger AUSSI la merge ECOLE globale C20:D79 si présente
    _unmerge_in_range(ws, first_row, last_row, 3, 18)

    # Vider toutes les valeurs cours
    for r in range(first_row, last_row + 1):
        for c in range(3, 19):
            cell = ws.cell(row=r, column=c)
            if not isinstance(cell, MergedCell):
                cell.value = None

    # Re-merger pattern uniforme
    for r in range(first_row, last_row + 1):
        ws.merge_cells(start_row=r, start_column=cols["promo"], end_row=r, end_column=cols["promo"] + 1)
        ws.merge_cells(start_row=r, start_column=cols["libelle"], end_row=r, end_column=cols["libelle"] + 3)
        ws.merge_cells(start_row=r, start_column=cols["date"], end_row=r, end_column=cols["date"] + 1)
        ws.merge_cells(start_row=r, start_column=cols["heures"], end_row=r, end_column=cols["heures"] + 1)
        ws.merge_cells(start_row=r, start_column=cols["taux"], end_row=r, end_column=cols["taux"] + 1)
        ws.merge_cells(start_row=r, start_column=cols["total"], end_row=r, end_column=cols["total"] + 1)

    # ECOLE merge global C:D sur toutes les rows utilisées (et pas plus)
    if n_sessions > 0:
        ws.merge_cells(
            start_row=first_row,
            start_column=cols["ecole"],
            end_row=last_session_row,
            end_column=cols["ecole"] + 1,
        )

    # Placer chaque session
    for i, (promo, libelle, d, heures, taux) in enumerate(flat):
        r = first_row + i

        # Date col K (mergée K:L)
        cd = ws.cell(row=r, column=cols["date"])
        cd.value = d
        cd.number_format = "DD/MM/YYYY"
        cd.font = Font(size=14)
        cd.alignment = Alignment(horizontal="center", vertical="center")

        # Heures col M (mergée M:N)
        ch = ws.cell(row=r, column=cols["heures"])
        ch.value = heures
        ch.font = Font(size=14)
        ch.alignment = Alignment(horizontal="center", vertical="center")

        # Taux col O (mergée O:P)
        ct = ws.cell(row=r, column=cols["taux"])
        ct.value = taux
        ct.font = Font(size=14)
        ct.alignment = Alignment(horizontal="center", vertical="center")

        # Total col Q (mergée Q:R) — formule
        col_taux_letter = chr(64 + cols["taux"])
        col_heures_letter = chr(64 + cols["heures"])
        cq = ws.cell(row=r, column=cols["total"])
        cq.value = f"={col_taux_letter}{r}*{col_heures_letter}{r}"
        cq.font = Font(size=14)
        cq.alignment = Alignment(horizontal="center", vertical="center")

    # ECOLE valeur en C20
    if n_sessions > 0:
        ce = ws.cell(row=first_row, column=cols["ecole"])
        ce.value = ecole_name
        ce.font = Font(bold=True, size=18)
        ce.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    # Promotion + Libellé valeurs (avec merges multi-rows si group >= 2)
    # Note : on a fait les merges 1-row par défaut. Pour les groupes multi-rows,
    # on démerge les rows correspondantes et on re-merge multi-row.
    for promo, ps, pe in promo_ranges:
        if pe > ps:
            # Multi-row : démerger les rows individuelles, re-merger multi-row
            for r in range(ps, pe + 1):
                for mr in list(ws.merged_cells.ranges):
                    if (
                        mr.min_row == r
                        and mr.max_row == r
                        and mr.min_col == cols["promo"]
                        and mr.max_col == cols["promo"] + 1
                    ):
                        ws.unmerge_cells(str(mr))
                        break
            ws.merge_cells(
                start_row=ps,
                start_column=cols["promo"],
                end_row=pe,
                end_column=cols["promo"] + 1,
            )
        cell = ws.cell(row=ps, column=cols["promo"])
        cell.value = promo
        cell.font = Font(size=16)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    for libelle, ls, le in libelle_subranges:
        if le > ls:
            for r in range(ls, le + 1):
                for mr in list(ws.merged_cells.ranges):
                    if (
                        mr.min_row == r
                        and mr.max_row == r
                        and mr.min_col == cols["libelle"]
                        and mr.max_col == cols["libelle"] + 3
                    ):
                        ws.unmerge_cells(str(mr))
                        break
            ws.merge_cells(
                start_row=ls,
                start_column=cols["libelle"],
                end_row=le,
                end_column=cols["libelle"] + 3,
            )
        cell = ws.cell(row=ls, column=cols["libelle"])
        cell.value = libelle
        cell.font = Font(size=14)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    # Masquer rows excédentaires (jamais delete)
    for r in range(last_session_row + 1, last_row + 1):
        ws.row_dimensions[r].hidden = True

    if masquer_jour:
        for c in (cols["date"], cols["date"] + 1):
            ws.column_dimensions[get_column_letter(c)].hidden = True


def apply_print_setup(ws: Worksheet, entite: str = "ROSARI"):
    """Fige la mise en page : portrait, ajusté à 1 page de large.

    Le template fait 18 colonnes ; sans ajustement, l'export PDF coupe les
    colonnes Taux/Total sur une 2e page.

    Args:
        ws: worksheet cible.
        entite: "ROSARI" ou "KAUR".
    """
    layout = _layout_for(entite)
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    cols = layout["cols_table"]
    derniere_col = get_column_letter(cols["total"] + 1)
    ws.print_area = f"A1:{derniere_col}{layout['derniere_ligne_utile']}"


def apply_rib(ws: Worksheet, entite: str = "ROSARI"):
    """Re-écrit IBAN/BIC dans les cells du template (idempotent).

    Args:
        ws: worksheet cible.
        entite: "ROSARI" ou "KAUR".
    """
    layout = _layout_for(entite)
    rib = _rib_for(entite)
    ws[layout["rib_iban_cell"]] = rib["IBAN"]
    ws[layout["rib_bic_cell"]] = rib["BIC"]


def post_mortem(ws: Worksheet, entite: str = "ROSARI") -> dict:
    """Retourne les cellules clés pour vérification après écriture.

    Args:
        ws: worksheet à inspecter.
        entite: "ROSARI" ou "KAUR".

    Returns:
        dict avec : titre, emetteur_nom, destinataire_nom, n_facture, siret,
                    ecole, premier_libelle, total_formule, iban, bic.
    """
    layout = _layout_for(entite)
    cols = layout["cols_table"]
    rmin_em, _ = layout["emetteur_rows"]
    rmin_dest, _ = layout["destinataire_rows"]

    return {
        "titre": ws.cell(row=layout["bandeau_row"], column=2).value,
        "emetteur_nom": ws.cell(row=rmin_em, column=layout["emetteur_col"]).value,
        "destinataire_nom": ws.cell(
            row=rmin_dest, column=layout["destinataire_col"]
        ).value,
        "n_facture": ws[layout["n_facture_cell"]].value,
        "siret": ws[layout["siret_cell"]].value,
        "ecole": ws.cell(row=layout["cours_first_row"], column=cols["ecole"]).value,
        "premier_libelle": ws.cell(
            row=layout["cours_first_row"], column=cols["libelle"]
        ).value,
        "total_formule": ws.cell(
            row=layout["total_row"], column=cols["total"]
        ).value,
        "iban": ws[layout["rib_iban_cell"]].value,
        "bic": ws[layout["rib_bic_cell"]].value,
    }
