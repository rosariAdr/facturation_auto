"""verification.py — Contrôle d'une installation à partir d'un profil (F-02).

`facturation profil verifier` passe ces contrôles avant tout lancement : un
problème bloquant (ERREUR) empêcherait de préparer les factures, un point
d'ATTENTION ne bloque pas encore mais le deviendra à un lot suivant.

Rien n'est écrit : la vérification ne fait que lire.
"""

from __future__ import annotations

import shutil
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from facturation import pcf_format as fmt
from facturation.config import Profil

OK, ATTENTION, ERREUR = "OK", "ATTENTION", "ERREUR"

ONGLETS_OBLIGATOIRES = ("Coordonnées", "Sessions")
ONGLETS_SUIVI = ("Cours", "Tarifs", "Écoles", "Devis")

_EMPLACEMENTS_LIBREOFFICE_WINDOWS = (
    Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
    Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
)


@dataclass(frozen=True)
class Constat:
    niveau: str
    sujet: str
    message: str


def trouver_libreoffice() -> Path | None:
    """Chemin de l'exécutable LibreOffice, ou None s'il est introuvable."""
    for nom in ("soffice", "libreoffice"):
        trouve = shutil.which(nom)
        if trouve:
            return Path(trouve)
    if sys.platform == "win32":
        for chemin in _EMPLACEMENTS_LIBREOFFICE_WINDOWS:
            if chemin.exists():
                return chemin
    return None


def verifier(profil: Profil) -> list[Constat]:
    """Contrôle le profil, les données et les outils. Renvoie la liste des constats."""
    constats: list[Constat] = [
        Constat(OK, "Profil", f"{profil.chemin} — {profil.nom}, entité {profil.entite}"),
    ]
    constats += _verifier_donnees(profil)
    constats += _verifier_gabarit(profil)
    constats += _verifier_agenda(profil)
    constats += _verifier_pdf(profil)
    return constats


def _verifier_donnees(profil: Profil) -> list[Constat]:
    if not profil.racine.is_dir():
        return [Constat(
            ERREUR, "Données",
            f"dossier introuvable : {profil.racine}. Si c'est un lecteur Google Drive "
            "pour ordinateur, vérifiez que l'application est lancée et la session ouverte.",
        )]
    constats = [Constat(OK, "Données", f"dossier trouvé : {profil.racine}")]

    if not profil.master.is_file():
        constats.append(Constat(ERREUR, "MASTER", f"fichier introuvable : {profil.master}"))
        return constats

    verrou = profil.master.with_name(f"~${profil.master.name}")
    if verrou.exists():
        constats.append(Constat(
            ATTENTION, "MASTER",
            "le fichier semble ouvert dans Excel : fermez-le avant toute écriture",
        ))

    try:
        wb = load_workbook(profil.master, read_only=True)
        onglets = set(wb.sheetnames)
        wb.close()
    except Exception as exc:  # fichier corrompu, verrouillé, non synchronisé…
        constats.append(Constat(ERREUR, "MASTER", f"illisible : {exc}"))
        return constats

    manquants = [o for o in ONGLETS_OBLIGATOIRES if o not in onglets]
    if manquants:
        constats.append(Constat(ERREUR, "MASTER", f"onglets obligatoires absents : {', '.join(manquants)}"))
        return constats
    absents_suivi = [o for o in ONGLETS_SUIVI if o not in onglets]
    if absents_suivi:
        constats.append(Constat(ATTENTION, "MASTER", f"onglets de suivi absents : {', '.join(absents_suivi)}"))
    constats.append(Constat(OK, "MASTER", f"{profil.master.name} lisible, {len(onglets)} onglets"))

    fmt.set_data_root(profil.racine)
    try:
        donnees = fmt._load_master_data(profil.entite, refresh=True)
    except RuntimeError as exc:
        constats.append(Constat(ERREUR, "Coordonnées", str(exc)))
        return constats
    constats.append(Constat(
        OK, "Coordonnées",
        f"émetteur « {donnees['emetteur']['nom']} », SIRET, IBAN et BIC présents",
    ))

    incomplets = sorted(code for code, dest in donnees["destinataires"].items() if dest is None)
    if incomplets:
        constats.append(Constat(
            ATTENTION, "Coordonnées",
            f"destinataires sans coordonnées (factures bloquées) : {', '.join(incomplets)}",
        ))
    return constats


def _verifier_gabarit(profil: Profil) -> list[Constat]:
    gabarit = fmt.gabarit_pour(profil.entite)
    if not gabarit.is_file():
        return [Constat(ERREUR, "Gabarit", f"introuvable : {gabarit}")]
    sorte = "standard" if gabarit == fmt.GABARIT_STANDARD else "propre à l'entité"
    return [Constat(OK, "Gabarit", f"{gabarit.name} ({sorte})")]


def _verifier_agenda(profil: Profil) -> list[Constat]:
    if profil.agenda_source == "api":
        return [Constat(ATTENTION, "Agenda", "source « api » : prévue en option, pas encore disponible")]
    fichier = profil.agenda_fichier_url_secrete
    if fichier is None:
        return [Constat(
            ATTENTION, "Agenda",
            "aucun fichier d'adresse iCal secrète déclaré ([agenda] fichier_url_secrete) : "
            "nécessaire à partir du lot 1",
        )]
    if not fichier.is_file():
        return [Constat(ATTENTION, "Agenda", f"fichier d'adresse iCal introuvable : {fichier}")]
    adresse = fichier.read_text(encoding="utf-8").strip()
    # L'adresse est un secret : on ne l'affiche jamais, on contrôle seulement sa forme.
    if not (adresse.startswith("https://") and adresse.endswith(".ics")):
        return [Constat(
            ATTENTION, "Agenda",
            f"le contenu de {fichier.name} ne ressemble pas à une adresse iCal (https://….ics)",
        )]
    return [Constat(OK, "Agenda", f"adresse iCal secrète présente dans {fichier.name} (non affichée)")]


def _verifier_pdf(profil: Profil) -> list[Constat]:
    if profil.pdf_moteur == "excel":
        return [Constat(ATTENTION, "PDF", "moteur « excel » : prévu en option (ADR021), pas encore disponible")]
    soffice = trouver_libreoffice()
    if soffice is None:
        return [Constat(ATTENTION, "PDF", "LibreOffice introuvable : nécessaire pour exporter les PDF")]
    return [Constat(OK, "PDF", f"LibreOffice trouvé : {soffice}")]
