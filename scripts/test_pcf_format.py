"""Tests de non-régression de pcf_format.py.

Nécessite un MASTER accessible (cf. racine de données). Les fichiers produits
vont dans un dossier temporaire, sauf si --garder est passé.

Usage :
    python scripts/test_pcf_format.py --data-root "/chemin/Factures"
    python scripts/test_pcf_format.py --garder ./sortie_tests
"""
import argparse
import sys
import tempfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import pcf_format as fmt

SORTIE = Path(tempfile.gettempdir()) / "pcf_tests"


def _new_wb() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def test_1_clone():
    """Le gabarit versionné est vide de coordonnées ; la structure est intacte."""
    print("\n══ TEST 1 — Clone du gabarit ══")
    wb = _new_wb()
    ws = fmt.clone_template(wb, "TEST_CLONE", entite="ROSARI")
    wb.save(SORTIE / "test1_clone.xlsx")

    pm = fmt.post_mortem(ws, entite="ROSARI")
    print(f"  Total formule : {pm['total_formule']!r}")
    print(f"  Émetteur      : {pm['emetteur_nom']!r} (attendu None — gabarit neutre)")
    print(f"  IBAN          : {pm['iban']!r} (attendu None — gabarit neutre)")

    assert pm["total_formule"] == "=SUM(Q20:R79)", pm["total_formule"]
    assert pm["emetteur_nom"] is None, pm["emetteur_nom"]
    assert pm["iban"] is None, pm["iban"]
    assert pm["destinataire_nom"] is None, pm["destinataire_nom"]
    print("  ✅ OK")


def test_2_coordonnees_injectees():
    """Émetteur, SIRET, IBAN et BIC sont injectés au runtime depuis MASTER."""
    print("\n══ TEST 2 — Injection depuis MASTER ══")
    wb = _new_wb()
    ws = fmt.clone_template(wb, "TEST_COORDS", entite="ROSARI")
    fmt.apply_emetteur(ws, entite="ROSARI")
    fmt.apply_metadata(
        ws, n_facture="R_TEST002", date_emission=date(2026, 4, 30), entite="ROSARI"
    )
    fmt.apply_rib(ws, entite="ROSARI")
    wb.save(SORTIE / "test2_coords.xlsx")

    pm = fmt.post_mortem(ws, entite="ROSARI")
    for cle in ("emetteur_nom", "siret", "iban", "bic"):
        print(f"  {cle:<14} : {pm[cle]!r}")
        assert pm[cle], f"{cle} non injecté"
    assert ws["C16"].value == "N° SIRET :", ws["C16"].value
    assert ws["C17"].value == "N° de Facture :", ws["C17"].value
    print("  ✅ OK")


def test_3_destinataire_complet():
    """La ligne rue ne doit pas disparaître dans la fusion M7:P8."""
    print("\n══ TEST 3 — Destinataire complet ══")
    wb = _new_wb()
    ws = fmt.clone_template(wb, "TEST_DEST", entite="ROSARI")
    dest = {
        "nom": "ECOLE TEST",
        "adresse_l1": "1 rue de la Paix",
        "adresse_l2": "75002 Paris",
        "adresse_l3": "France",
        "email": "compta@test.fr",
    }
    fmt.apply_destinataire(ws, dest=dest, entite="ROSARI")
    wb.save(SORTIE / "test3_dest.xlsx")

    lues = [ws.cell(row=r, column=13).value for r in range(7, 11)]
    print(f"  M7:M10 = {lues}")
    assert lues[0] == "ECOLE TEST", lues
    assert lues[1] == "1 rue de la Paix", "rue perdue dans la fusion M7:P8"
    assert lues[2] == "75002 Paris, France", lues
    assert lues[3] == "compta@test.fr", lues
    print("  ✅ OK")


def test_4_destinataire_vide():
    """dest=None laisse l'encadré mais aucune valeur."""
    print("\n══ TEST 4 — Destinataire vide ══")
    wb = _new_wb()
    ws = fmt.clone_template(wb, "TEST_DEST_VIDE", entite="ROSARI")
    fmt.apply_destinataire(ws, dest=None, entite="ROSARI")
    wb.save(SORTIE / "test4_dest_vide.xlsx")

    for r in range(7, 11):
        assert ws.cell(row=r, column=13).value is None, f"M{r} non vide"
    print("  ✅ OK")


def test_5_une_session():
    """Une session : row 20 remplie, 21-79 masquées, formule correcte."""
    print("\n══ TEST 5 — Une seule session ══")
    wb = _new_wb()
    ws = fmt.clone_template(wb, "TEST_1_SESSION", entite="ROSARI")
    fmt.apply_metadata(
        ws, n_facture="R_TEST005", date_emission=date(2026, 4, 30), entite="ROSARI"
    )
    groups = [("BTS SIO A1", "Maths BTS SIO", [(date(2026, 3, 12), 3.5, 60)])]
    fmt.place_sessions(
        ws, ecole_name="DOR (NEXA Digital School)", groups=groups, entite="ROSARI"
    )
    wb.save(SORTIE / "test5_1session.xlsx")

    pm = fmt.post_mortem(ws, entite="ROSARI")
    assert pm["n_facture"] == "R_TEST005", pm["n_facture"]
    assert pm["ecole"] == "DOR (NEXA Digital School)", pm["ecole"]
    assert pm["premier_libelle"] == "Maths BTS SIO", pm["premier_libelle"]
    assert "30/04/2026" in str(pm["titre"]), pm["titre"]

    n_hidden = sum(1 for r in range(21, 80) if ws.row_dimensions[r].hidden)
    print(f"  Rows masquées 21-79 : {n_hidden}/59")
    assert n_hidden == 59, n_hidden
    assert ws.cell(row=20, column=13).value == 3.5
    assert ws.cell(row=20, column=15).value == 60
    assert ws.cell(row=20, column=17).value == "=O20*M20"
    print("  ✅ OK")


def test_6_bon_commande_et_jour_masque():
    """Ligne Bon de commande écrite, colonnes Jour masquées à la demande."""
    print("\n══ TEST 6 — Bon de commande + Jour masqué ══")
    wb = _new_wb()
    ws = fmt.clone_template(wb, "TEST_BC", entite="ROSARI")
    fmt.apply_metadata(
        ws,
        n_facture="R_TEST006",
        date_emission=date(2026, 7, 31),
        entite="ROSARI",
        ref_bon_commande="PO-TEST-0001",
    )
    fmt.place_sessions(
        ws,
        ecole_name="EPITA",
        groups=[("AIS S1", "UE de test", [(date(2026, 7, 31), 18, 90)])],
        entite="ROSARI",
        masquer_jour=True,
    )
    fmt.apply_print_setup(ws, entite="ROSARI")
    wb.save(SORTIE / "test6_bon_commande.xlsx")

    print(f"  C14/D14 : {ws['C14'].value!r} {ws['D14'].value!r}")
    assert ws["C14"].value == "Bon de commande :", ws["C14"].value
    assert ws["D14"].value == "PO-TEST-0001", ws["D14"].value
    assert ws.column_dimensions["K"].hidden, "colonne K (Jour) non masquée"
    assert ws.column_dimensions["L"].hidden, "colonne L (Jour) non masquée"
    assert ws.page_setup.orientation == "portrait"
    assert ws.print_area, "print_area non définie"
    print(f"  print_area : {ws.print_area}")
    print("  ✅ OK")


TESTS = [
    test_1_clone,
    test_2_coordonnees_injectees,
    test_3_destinataire_complet,
    test_4_destinataire_vide,
    test_5_une_session,
    test_6_bon_commande_et_jour_masque,
]


def main():
    global SORTIE
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--data-root", type=Path, default=None)
    parser.add_argument(
        "--garder", type=Path, default=None,
        help="Dossier où conserver les XLSX produits (défaut : dossier temporaire)",
    )
    args = parser.parse_args()

    if args.data_root:
        fmt.set_data_root(args.data_root)
    if args.garder:
        SORTIE = args.garder
    SORTIE.mkdir(parents=True, exist_ok=True)

    print(f"Racine de données : {fmt.get_data_root()}")
    print(f"Sorties           : {SORTIE}")
    for test in TESTS:
        test()
    print(f"\n✅ {len(TESTS)}/{len(TESTS)} tests passés")


if __name__ == "__main__":
    main()
