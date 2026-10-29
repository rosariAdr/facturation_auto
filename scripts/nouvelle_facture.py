"""nouvelle_facture.py — Génère une facture depuis une spec JSON.

Remplace les scripts one-shot : la spec décrit le quoi, ce script applique le PCF
étape 5 (format) et étape 7 (écriture + post-mortem). Les étapes humaines du
protocole — périmètre, coordonnées, aperçu — restent à la charge de l'appelant
(cf. references/PROC00_pcf.md).

Usage :
    python scripts/nouvelle_facture.py spec.json
    python scripts/nouvelle_facture.py spec.json --data-root "/chemin/Factures"
    python scripts/nouvelle_facture.py spec.json --dry-run    # aperçu seul

Format de la spec (voir examples/spec_exemple.json) :
{
  "entite": "ROSARI",
  "ecole": "EPITA",
  "ecole_label": "EPITA",
  "fichier_sortie": "ROSARI/EPITA/ROSARI_EPITA.xlsx",
  "onglet": "R_EPITA010 Sept 2026",
  "n_facture": "R_EPITA010",
  "date_emission": "2026-09-30",
  "ref_bon_commande": "PO-26-ASSOTA-02001",
  "masquer_jour": true,
  "lignes": [
    {"promotion": "AIS S1", "libelle": "[2026][PI][AIS S1][...]",
     "seances": [{"date": "2026-09-15", "heures": 18, "taux": 90}]}
  ]
}
"""
import argparse
import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

try:
    from facturation import config, pcf_format as fmt
except ModuleNotFoundError:  # exécution depuis le dépôt sans installation
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    from facturation import config, pcf_format as fmt


def _parse_date(valeur) -> date:
    """Accepte 'YYYY-MM-DD' ou un objet date déjà construit."""
    if isinstance(valeur, date):
        return valeur
    return datetime.strptime(str(valeur), "%Y-%m-%d").date()


def charger_spec(chemin: Path) -> dict:
    """Lit la spec JSON et vérifie que les clés obligatoires sont présentes."""
    spec = json.loads(chemin.read_text(encoding="utf-8"))
    obligatoires = {
        "entite", "ecole", "fichier_sortie", "onglet",
        "n_facture", "date_emission", "lignes",
    }
    manquantes = obligatoires - set(spec)
    if manquantes:
        raise ValueError(f"Clés manquantes dans {chemin.name} : {sorted(manquantes)}")
    if not spec["lignes"]:
        raise ValueError("La spec ne contient aucune ligne de cours.")
    return spec


def construire_groupes(lignes: list) -> list:
    """Transforme les lignes de la spec en `groups` pour place_sessions()."""
    groupes = []
    for ligne in lignes:
        seances = [
            (_parse_date(s["date"]), float(s["heures"]), float(s["taux"]))
            for s in ligne["seances"]
        ]
        groupes.append((ligne["promotion"], ligne["libelle"], seances))
    return groupes


def apercu(spec: dict, groupes: list, dest: dict) -> str:
    """Construit l'aperçu du PCF étape 6 (gate humain avant écriture)."""
    lignes_txt = []
    total_h = total_eur = 0.0
    for promo, libelle, seances in groupes:
        for d, heures, taux in seances:
            total_h += heures
            total_eur += heures * taux
            libelle_court = libelle.replace("\n", " ")[:45]
            lignes_txt.append(
                f"  | {d.strftime('%d/%m/%Y')} | {promo:<12} | {libelle_court:<45} "
                f"| {heures:>5}h | {taux:>6.2f} € | {heures * taux:>9.2f} €"
            )
    bc = spec.get("ref_bon_commande") or "—"
    return "\n".join([
        "====== APERÇU FACTURE ======",
        f"Entité        : {spec['entite']}",
        f"N°            : {spec['n_facture']}",
        f"Date émission : {_parse_date(spec['date_emission']).strftime('%d/%m/%Y')}",
        f"Destinataire  : {dest['nom']} — {dest.get('adresse_l1', '')}, "
        f"{dest.get('adresse_l2', '')}",
        f"Bon de commande : {bc}",
        f"Fichier       : {spec['fichier_sortie']} > onglet '{spec['onglet']}'",
        "",
        "DÉTAIL DES SÉANCES :",
        *lignes_txt,
        "",
        f"TOTAL HEURES : {total_h}h",
        f"TOTAL HT     : {total_eur:.2f} €",
        "============================",
    ])


def generer(spec: dict, dry_run: bool = False) -> dict:
    """Applique le PCF étapes 5 et 7. Retourne le post-mortem (dict vide si dry-run)."""
    entite = spec["entite"]
    dest = fmt.get_destinataire(spec["ecole"], entite=entite)
    if dest is None:
        raise RuntimeError(
            f"Destinataire '{spec['ecole']}' absent de MASTER.Coordonnées — "
            "gate étape 4 du PCF : compléter avant de générer "
            "(cf. references/ONBOARDING.md)."
        )

    groupes = construire_groupes(spec["lignes"])
    print(apercu(spec, groupes, dest))
    if dry_run:
        print("\n[dry-run] Aucun fichier écrit.")
        return {}

    sortie = fmt.get_data_root() / spec["fichier_sortie"]
    sortie.parent.mkdir(parents=True, exist_ok=True)
    if sortie.exists():
        wb = load_workbook(sortie)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    onglet = spec["onglet"]
    if onglet in wb.sheetnames:
        del wb[onglet]
        print(f"\nAncien onglet '{onglet}' remplacé.")

    ws = fmt.clone_template(wb, onglet, entite=entite)
    fmt.apply_emetteur(ws, entite=entite)
    fmt.apply_destinataire(ws, dest=dest, entite=entite)
    fmt.apply_metadata(
        ws,
        n_facture=spec["n_facture"],
        date_emission=_parse_date(spec["date_emission"]),
        entite=entite,
        ref_bon_commande=spec.get("ref_bon_commande"),
    )
    fmt.place_sessions(
        ws,
        ecole_name=spec.get("ecole_label", spec["ecole"]),
        groups=groupes,
        entite=entite,
        masquer_jour=bool(spec.get("masquer_jour", False)),
    )
    fmt.apply_rib(ws, entite=entite)
    fmt.apply_print_setup(ws, entite=entite)
    wb.save(sortie)
    print(f"\nÉcrit : {sortie}")

    pm = fmt.post_mortem(ws, entite=entite)
    print("\n── Post-mortem (PCF étape 7.2) ──")
    for cle, valeur in pm.items():
        print(f"  {cle:<18} : {valeur!r}")
    return pm


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("spec", type=Path, help="Fichier JSON de spec")
    parser.add_argument(
        "--data-root",
        type=Path,
        default=None,
        help="Racine des données métier (défaut : profil, $PCF_DATA_ROOT ou dossier courant)",
    )
    parser.add_argument(
        "--profil",
        type=Path,
        default=None,
        help="Profil utilisateur (défaut : profil standard s'il existe)",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Affiche l'aperçu (PCF étape 6) sans rien écrire",
    )
    args = parser.parse_args()

    config.resoudre_racine(args.data_root, args.profil)
    print(f"Racine de données : {fmt.get_data_root()}")

    generer(charger_spec(args.spec), dry_run=args.dry_run)


if __name__ == "__main__":
    main()
