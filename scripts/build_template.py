"""build_template.py — Reconstruit le gabarit canonique ROSARI depuis une facture.

Source : une facture existante servant de référence visuelle (READ-ONLY).
Sortie : facturation/assets/PCF_facture_standard.xlsx > MODELE.

Script de maintenance : à relancer seulement pour refondre le gabarit. En usage
courant, le gabarit livré dans facturation/assets/ suffit.

Modifications apportées vs la facture source :
- Étend la zone cours de 7 → 60 rows (insert 53 rows entre row 26 et row 27)
- Préserve manuellement les merges des rows 27+ (shift +53) — openpyxl insert_rows
  ne shift pas les merges proprement
- Préserve les hauteurs de ligne des rows 27+ (shift +53)
- Ajoute encadrés bordures fines noires sur émetteur (C7:E11) et destinataire (M7:P10)
- Vide toutes les données : cours, destinataire, N° facture, ET les coordonnées
  émetteur / SIRET / RIB. Elles sont injectées au runtime depuis MASTER.Coordonnées
  par apply_emetteur / apply_metadata / apply_rib — rien de personnel ne doit être
  figé dans un gabarit versionné.
- Met à jour formules : Q80 = =SUM(Q20:R79), E83 = =Q80

Usage :
    python scripts/build_template.py --source "<...>/ROSARI_DOR.xlsx" --onglet "DOR001 012026"
"""
import argparse
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Border, Side

PACKAGE_ROOT = Path(__file__).resolve().parent.parent
DST_FILE = PACKAGE_ROOT / "facturation" / "assets" / "PCF_facture_standard.xlsx"

# Renseignées par main() depuis les arguments de la ligne de commande
SRC_FILE: Path
SRC_SHEET: str

# Cellules vidées avant sauvegarde : réinjectées au runtime depuis MASTER
CELLULES_NEUTRALISEES = ["C7", "C8", "C9", "C10", "C11", "D16", "D17", "D88", "D89"]

THIN = Side(style="thin", color="000000")
SHIFT = 53  # nb de rows insérées entre row 26 et l'ancien row 27


def deep_copy_sheet(src_ws, dst_ws):
    """Cell-by-cell deep copy : valeurs + styles + merges + dimensions."""
    for row in src_ws.iter_rows():
        for cell in row:
            if cell.value is not None or cell.has_style:
                new_cell = dst_ws[cell.coordinate]
                new_cell.value = cell.value
                if cell.has_style:
                    new_cell.font = copy(cell.font)
                    new_cell.fill = copy(cell.fill)
                    new_cell.border = copy(cell.border)
                    new_cell.alignment = copy(cell.alignment)
                    new_cell.number_format = cell.number_format
    for col_letter, dim in src_ws.column_dimensions.items():
        if dim.width:
            dst_ws.column_dimensions[col_letter].width = dim.width
    for r, dim in src_ws.row_dimensions.items():
        if dim.height:
            dst_ws.row_dimensions[r].height = dim.height
    for mr in src_ws.merged_cells.ranges:
        try:
            dst_ws.merge_cells(str(mr))
        except Exception as exc:
            print(f"  ⚠ merge {mr} failed : {exc}")


def unmerge_in_zone(ws, rmin, rmax, cmin=1, cmax=30):
    """Démerge toutes les plages incluses dans la zone."""
    to_remove = []
    for mr in list(ws.merged_cells.ranges):
        if (
            mr.min_row >= rmin
            and mr.max_row <= rmax
            and mr.min_col >= cmin
            and mr.max_col <= cmax
        ):
            to_remove.append(str(mr))
    for mr_str in to_remove:
        ws.unmerge_cells(mr_str)
    return len(to_remove)


def clear_values(ws, rmin, rmax, cmin=1, cmax=30):
    """Vide les valeurs des cellules non-merged dans la zone."""
    for r in range(rmin, rmax + 1):
        for c in range(cmin, cmax + 1):
            cell = ws.cell(row=r, column=c)
            if not isinstance(cell, MergedCell):
                cell.value = None


def apply_box_borders(ws, rmin, rmax, cmin, cmax):
    """Encadré : top/bottom sur extrémités, left/right sur extrémités."""
    for r in range(rmin, rmax + 1):
        for c in range(cmin, cmax + 1):
            cell = ws.cell(row=r, column=c)
            if isinstance(cell, MergedCell):
                continue
            existing = cell.border
            new_top = THIN if r == rmin else existing.top
            new_bottom = THIN if r == rmax else existing.bottom
            new_left = THIN if c == cmin else existing.left
            new_right = THIN if c == cmax else existing.right
            cell.border = Border(
                top=new_top, bottom=new_bottom, left=new_left, right=new_right
            )


def apply_grid_borders(ws, rmin, rmax, cmin, cmax):
    """Grid borders sur toutes les cellules (4 côtés)."""
    for r in range(rmin, rmax + 1):
        for c in range(cmin, cmax + 1):
            cell = ws.cell(row=r, column=c)
            if isinstance(cell, MergedCell):
                continue
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def main():
    print(f"Source : {SRC_FILE}")
    print(f"Cible  : {DST_FILE}")

    src_wb = load_workbook(SRC_FILE)
    src_ws = src_wb[SRC_SHEET]
    print(f"  Source dim : {src_ws.max_row}r × {src_ws.max_column}c")

    # ── 1. Sauvegarder les merges et heights des rows 27+ AVANT toute modif ──
    bottom_merges = []
    bottom_heights = {}
    for mr in src_ws.merged_cells.ranges:
        if mr.min_row >= 27:
            bottom_merges.append((mr.min_row, mr.max_row, mr.min_col, mr.max_col))
    for r, dim in src_ws.row_dimensions.items():
        if r >= 27 and dim.height:
            bottom_heights[r] = dim.height
    print(f"  Sauvegardé {len(bottom_merges)} merges bottom + {len(bottom_heights)} heights bottom")

    # ── 2. Deep copy ──
    dst_wb = Workbook()
    dst_wb.remove(dst_wb.active)
    dst_ws = dst_wb.create_sheet("MODELE")
    deep_copy_sheet(src_ws, dst_ws)
    print(f"  Deep copy effectué : {dst_ws.max_row}r × {dst_ws.max_column}c")

    # ── 3. Démerger zone cours (rows 20-26) ──
    n_unmerged = unmerge_in_zone(dst_ws, 20, 26, 3, 18)
    print(f"  Démergé {n_unmerged} plages cours zone (rows 20-26)")
    clear_values(dst_ws, 20, 26, 3, 18)

    # ── 4. Démerger TOUTES les plages dans rows 27+ avant insert_rows ──
    # (insert_rows ne les shift pas correctement, on les recrée manuellement)
    bottom_to_remove = [
        str(mr) for mr in list(dst_ws.merged_cells.ranges) if mr.min_row >= 27
    ]
    for mr_str in bottom_to_remove:
        dst_ws.unmerge_cells(mr_str)
    print(f"  Démergé {len(bottom_to_remove)} plages bottom (rows 27+) avant shift")

    # ── 5. Insert 53 rows ──
    dst_ws.insert_rows(27, SHIFT)
    print(f"  Inséré {SHIFT} rows entre row 26 et l'ancien row 27 → row {27 + SHIFT}")

    # ── 6. Recréer les merges bottom shiftées de +SHIFT ──
    n_shifted = 0
    for min_row, max_row, min_col, max_col in bottom_merges:
        new_min_row = min_row + SHIFT
        new_max_row = max_row + SHIFT
        try:
            dst_ws.merge_cells(
                start_row=new_min_row,
                start_column=min_col,
                end_row=new_max_row,
                end_column=max_col,
            )
            n_shifted += 1
        except Exception as exc:
            print(f"  ⚠ merge shift {min_row}→{new_min_row} failed : {exc}")
    print(f"  Re-créé {n_shifted}/{len(bottom_merges)} merges bottom shiftées")

    # ── 7. Recréer les heights bottom shiftées ──
    for r, h in bottom_heights.items():
        new_r = r + SHIFT
        dst_ws.row_dimensions[new_r].height = h
    print(f"  Heights bottom : {len(bottom_heights)} rows shiftées")

    # ── 8. Pattern cours uniforme rows 20-79 ──
    for r in range(20, 80):
        dst_ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)   # E:F
        dst_ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=10)  # G:J
        dst_ws.merge_cells(start_row=r, start_column=11, end_row=r, end_column=12) # K:L
        dst_ws.merge_cells(start_row=r, start_column=13, end_row=r, end_column=14) # M:N
        dst_ws.merge_cells(start_row=r, start_column=15, end_row=r, end_column=16) # O:P
        dst_ws.merge_cells(start_row=r, start_column=17, end_row=r, end_column=18) # Q:R
        dst_ws.row_dimensions[r].height = 46.5

    dst_ws.merge_cells(start_row=20, start_column=3, end_row=79, end_column=4)
    print(f"  Pattern cours appliqué (rows 20-79, 60 rows)")

    apply_grid_borders(dst_ws, 20, 79, 3, 18)
    print(f"  Bordures grid rows 20-79")

    # ── 9. Encadrés émetteur (C7:E11) et destinataire (M7:P10) ──
    apply_box_borders(dst_ws, 7, 11, 3, 5)
    apply_box_borders(dst_ws, 7, 10, 13, 16)
    print(f"  Encadrés émetteur (C7:E11) + destinataire (M7:P10)")

    # ── 10. Vider destinataire M7-M10 ──
    for r in range(7, 11):
        cell = dst_ws.cell(row=r, column=13)
        if not isinstance(cell, MergedCell):
            cell.value = None

    # ── 11-12. Neutraliser toute donnée personnelle (émetteur, SIRET, RIB) ──
    # Le gabarit est versionné : il ne porte que la mise en forme. Les valeurs
    # viennent de MASTER.Coordonnées au runtime.
    for coord in CELLULES_NEUTRALISEES:
        dst_ws[coord] = None
    print(f"  Neutralisé {len(CELLULES_NEUTRALISEES)} cellules de coordonnées")

    # ── 13. Update formules ──
    dst_ws["Q80"] = "=SUM(Q20:R79)"
    dst_ws["E83"] = "=Q80"
    print(f"  Formules : Q80 = =SUM(Q20:R79), E83 = =Q80")

    # ── 13bis. Forcer hauteurs lignes pour lisibilité Excel + Google Sheets ──
    # Rows émetteur/destinataire 7-11 : suffisamment hautes pour le texte centré
    for r in range(7, 12):
        dst_ws.row_dimensions[r].height = 22.0
    # TOTAL row 80
    dst_ws.row_dimensions[80].height = 28.0
    dst_ws.row_dimensions[81].height = 14.25
    # À payer + TVA mentions (rows 83-86)
    dst_ws.row_dimensions[83].height = 28.0
    dst_ws.row_dimensions[84].height = 18.0
    dst_ws.row_dimensions[85].height = 18.0
    # RIB title + IBAN + BIC (rows 87-89)
    dst_ws.row_dimensions[87].height = 28.0
    dst_ws.row_dimensions[88].height = 28.0  # IBAN — texte long
    dst_ws.row_dimensions[89].height = 28.0  # BIC
    # Conditions paiement + pénalités (rows 91-93)
    dst_ws.row_dimensions[91].height = 22.0
    dst_ws.row_dimensions[92].height = 18.0
    dst_ws.row_dimensions[93].height = 30.0  # texte long pénalités
    print(f"  Hauteurs rows ajustées (émetteur/dest + bottom)")

    # ── 14. Sauver ──
    DST_FILE.parent.mkdir(parents=True, exist_ok=True)
    dst_wb.save(DST_FILE)
    print(f"\n✅ Template sauvegardé : {DST_FILE}")
    print(f"   Taille : {DST_FILE.stat().st_size} octets")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument(
        "--source", type=Path, required=True,
        help="Classeur contenant la facture de référence (lu seulement)",
    )
    parser.add_argument(
        "--onglet", required=True, help="Nom de l'onglet de référence dans --source",
    )
    args = parser.parse_args()
    SRC_FILE = args.source
    SRC_SHEET = args.onglet
    if not SRC_FILE.exists():
        raise SystemExit(f"Source introuvable : {SRC_FILE}")
    main()
