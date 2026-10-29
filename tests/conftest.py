"""Fixtures partagées : données fictives uniquement (aucune donnée réelle dans le dépôt).

`racine_fictive` construit dans un dossier temporaire une racine de données complète :
`ENTITE/ENTITE_MASTER.xlsx` avec émetteur, SIRET, IBAN et BIC factices, deux
destinataires renseignés et un destinataire vide.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from facturation import pcf_format as fmt

ENTETES = {
    "Sessions": [
        "Date_log", "École", "Période_facture", "N°_Facture", "Date_séance", "Début", "Fin",
        "Durée_h", "Promotion", "Module / UE", "Taux", "Total HT", "Summary calendrier verbatim",
    ],
    "Cours": [
        "École", "Module / Matière", "Classe / Groupe", "Volume contractuel (h)",
        "Heures dispensées (h)", "Heures facturées (h)", "Heures à facturer (h)",
        "N° Facture liée", "N° Devis lié", "Statut",
    ],
    "Tarifs": ["École", "Module / Matière", "Classe / Programme", "Taux HT/h", "Notes"],
    "Écoles": ["École", "Code facturation", "Entité légale", "Contact", "Email contact", "Statut mission"],
    "Devis": ["École", "Module", "N° Devis", "Date émission", "Montant HT", "Statut", "PO reçu", "N° Facture générée"],
}


def creer_master(racine: Path, entite: str = "ROSARI", cours: list[list] | None = None) -> Path:
    """Écrit un MASTER fictif sous `racine/entite/` et renvoie son chemin."""
    dossier = racine / entite
    dossier.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Coordonnées")
    ws.append(["Type", "Code", "Nom / Raison sociale", "Adresse L1", "Adresse L2", "Adresse L3", "Email", "Téléphone"])
    ws.append(["Émetteur", entite, f"FORMATEUR FICTIF {entite}", "1 rue de l'Exemple,", "95000 Villefictive", "", "formateur@example.org", "+33600000000"])
    ws.append(["Info", entite, "SIRET", "00000000000000"])
    ws.append(["Info", entite, "IBAN", "FR76 0000 0000 0000 0000 0000 000"])
    ws.append(["Info", entite, "BIC", "TESTFRPP"])
    ws.append(["Info", entite, "TVA", "TVA non applicable - art. 293B CGI"])
    ws.append(["Destinataire", "EPITA", "COMPTABILITE ECOLE A", "10 avenue Test", "94000 TESTVILLE", "France", "compta@ecole-a.example"])
    ws.append(["Destinataire", "DOR", "ECOLE B", "20 rue Test", "75000 Paris", "", "factures@ecole-b.example"])
    ws.append(["Destinataire", "VIDE", "", "", "", "", ""])
    for nom, entetes in ENTETES.items():
        feuille = wb.create_sheet(nom)
        feuille.append(entetes)
        if nom == "Cours":
            for ligne in cours or []:
                feuille.append(ligne)
    chemin = dossier / f"{entite}_MASTER.xlsx"
    wb.save(chemin)
    return chemin


@pytest.fixture
def racine_fictive(tmp_path: Path) -> Path:
    """Racine de données fictive, branchée sur `pcf_format` le temps du test."""
    racine = tmp_path / "donnees"
    creer_master(racine)
    ancienne = fmt._DATA_ROOT
    fmt.set_data_root(racine)
    yield racine
    fmt._DATA_ROOT = ancienne
    fmt._MASTER_DATA_CACHE.clear()
