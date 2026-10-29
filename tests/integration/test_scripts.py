"""Scripts et commande `facturation` sur données fictives (T-CLI-01, T-CLI-02, T-COH-01, T-VER-01)."""

from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from facturation import config
from tests.conftest import creer_master

pytestmark = pytest.mark.integration

DEPOT = Path(__file__).resolve().parents[2]
SPEC_EXEMPLE = DEPOT / "examples" / "spec_exemple.json"


def _lancer(*args: str, env_extra: dict | None = None) -> subprocess.CompletedProcess:
    env = {**os.environ, "PYTHONIOENCODING": "utf-8", **(env_extra or {})}
    env.pop("PCF_DATA_ROOT", None)
    return subprocess.run(
        [sys.executable, *args], capture_output=True, text=True, encoding="utf-8", env=env, cwd=DEPOT,
    )


def test_t_cli_01_apercu_dry_run(tmp_path):
    """T-CLI-01 — l'aperçu affiche les totaux et n'écrit aucun fichier."""
    racine = tmp_path / "donnees"
    creer_master(racine)
    res = _lancer("scripts/nouvelle_facture.py", str(SPEC_EXEMPLE), "--data-root", str(racine), "--dry-run")
    assert res.returncode == 0, res.stderr
    assert "TOTAL HT     : 1890.00 €" in res.stdout
    assert "[dry-run] Aucun fichier écrit." in res.stdout
    assert not (racine / "ROSARI" / "EPITA").exists()


def test_t_cli_02_ecriture_et_post_mortem(tmp_path):
    """T-CLI-02 — l'écriture produit l'onglet attendu et un post-mortem complet."""
    racine = tmp_path / "donnees"
    creer_master(racine)
    res = _lancer("scripts/nouvelle_facture.py", str(SPEC_EXEMPLE), "--data-root", str(racine))
    assert res.returncode == 0, res.stderr
    spec = json.loads(SPEC_EXEMPLE.read_text(encoding="utf-8"))
    wb = load_workbook(racine / spec["fichier_sortie"])
    ws = wb[spec["onglet"]]
    assert ws["D17"].value == spec["n_facture"]
    assert ws["D14"].value == spec["ref_bon_commande"]
    assert ws["M7"].value == "COMPTABILITE ECOLE A"
    post_mortem = {
        ligne.split(":")[0].strip(): ligne for ligne in res.stdout.splitlines() if ligne.startswith("  ")
    }
    for cle in ("destinataire_nom", "siret", "iban"):
        assert cle in post_mortem and "None" not in post_mortem[cle], post_mortem.get(cle)


def test_t_coh_01_ecart_volontaire_detecte(tmp_path):
    """T-COH-01 — un écart dispensées ≠ facturées + à facturer est signalé."""
    racine = tmp_path / "donnees"
    creer_master(racine, cours=[
        ["DOR", "Module juste", "A1", 0, 10, 6, 4, "R_DOR001", None, "En cours"],
        ["DOR", "Module faux", "A2", 0, 10, 6, 2, "R_DOR001", None, "En cours"],
    ])
    res = _lancer("scripts/check_coherence.py", "--entite", "ROSARI", "--data-root", str(racine))
    assert res.returncode == 0, res.stderr
    assert "Écarts" in res.stdout
    assert "Module faux" in res.stdout.split("Écarts")[1]
    assert "Module juste" not in res.stdout.split("Écarts")[1]


def test_t_ver_01_verifier_installation(tmp_path):
    """T-VER-01 — `facturation profil verifier` : 0 erreur sur une installation fictive complète."""
    racine = tmp_path / "donnees"
    creer_master(racine)
    profil = tmp_path / "profil.toml"
    profil.write_text(config.rediger("Personne A", "ROSARI", str(racine)), encoding="utf-8")
    res = _lancer("-m", "facturation", "profil", "verifier", "--profil", str(profil))
    assert res.returncode == 0, res.stdout + res.stderr
    assert "0 erreur(s)" in res.stdout
    assert "destinataires sans coordonnées (factures bloquées) : VIDE" in res.stdout


def test_t_ver_02_dossier_de_donnees_absent(tmp_path):
    """T-VER-02 — dossier de données introuvable : erreur bloquante, code de retour 1."""
    profil = tmp_path / "profil.toml"
    profil.write_text(config.rediger("Personne A", "ROSARI", str(tmp_path / "absent")), encoding="utf-8")
    res = _lancer("-m", "facturation", "profil", "verifier", "--profil", str(profil))
    assert res.returncode == 1
    assert "[ERREUR]" in res.stdout and "dossier introuvable" in res.stdout


def test_profil_init_refuse_d_ecraser(tmp_path):
    cible = tmp_path / "profil.toml"
    args = ["-m", "facturation", "profil", "init", "--nom", "Personne A", "--entite", "ROSARI",
            "--racine", str(tmp_path), "--emplacement", str(cible)]
    assert _lancer(*args).returncode == 0
    second = _lancer(*args)
    assert second.returncode == 1 and "--force" in second.stderr


ADRESSE_FICTIVE = "https://calendar.google.com/calendar/ical/demo%40example.org/private-0000/basic.ics"


def test_t_cfg_02_enregistrer_adresse_ical(tmp_path):
    """T-CFG-02 — `profil agenda` range l'adresse hors du profil et la déclare ; verifier la voit."""
    racine = tmp_path / "donnees"
    creer_master(racine)
    profil = tmp_path / "profil.toml"
    profil.write_text(config.rediger("Personne A", "ROSARI", str(racine)), encoding="utf-8")
    res = _lancer("-m", "facturation", "profil", "agenda", "--profil", str(profil), "--adresse", ADRESSE_FICTIVE)
    assert res.returncode == 0, res.stdout + res.stderr
    assert (tmp_path / "agenda.secret").read_text(encoding="utf-8").strip() == ADRESSE_FICTIVE
    assert ADRESSE_FICTIVE not in profil.read_text(encoding="utf-8")
    assert config.charger(profil).agenda_fichier_url_secrete == tmp_path / "agenda.secret"
    verif = _lancer("-m", "facturation", "profil", "verifier", "--profil", str(profil))
    assert "[OK]        Agenda" in verif.stdout
    assert ADRESSE_FICTIVE not in verif.stdout


def test_t_cfg_03_adresse_ical_invalide_refusee(tmp_path):
    """T-CFG-03 — une adresse qui n'est pas un lien .ics est refusée, rien n'est écrit."""
    profil = tmp_path / "profil.toml"
    profil.write_text(config.rediger("Personne A", "ROSARI", str(tmp_path)), encoding="utf-8")
    avant = profil.read_text(encoding="utf-8")
    res = _lancer("-m", "facturation", "profil", "agenda", "--profil", str(profil),
                  "--adresse", "https://calendar.google.com/calendar/u/0/r")
    assert res.returncode == 1 and "adresse iCal invalide" in res.stderr
    assert not (tmp_path / "agenda.secret").exists()
    assert profil.read_text(encoding="utf-8") == avant
